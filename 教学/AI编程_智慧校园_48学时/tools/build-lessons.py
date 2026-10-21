"""Build lesson Markdown/HTML/DOCX from one editable source; use bundled Python."""
from pathlib import Path
import json,html,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor,Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image
ROOT=Path(__file__).resolve().parent.parent; OUT=ROOT/'lesson-plans'
data=json.loads((OUT/'lesson-source.json').read_text()); lessons=data['lessons']
assert len(lessons)==12 and sum(x['hours'] for x in lessons)==48 and sum(x['minutes'] for x in lessons)==2160
intro='本教案用于带领学过一门编程语言、尚无完整项目经验的学生，使用通义等国内网页AI工具，逐步完成智慧校园教学平台。学生实现业务小块并验证结果，教师提供认证、连接、事务和故障实验基础。每项实现都对应需求、代码与证据。'
arr='建议每学时45分钟，每次4学时，共12次课；建议30人，两人一组。每次教学时间180分钟，课间休息另计。学校可按实际课表调整节次，保留48学时总量及核心验收。'
md=['# '+data['title'],intro,arr,'## 教学实施与评价','每组轮换操作与审查角色，每学时至少交换一次。网页AI不可用时使用离线练习；敏感配置只留本机。UI允许自主设计，AIRIOT为0必修学时的只读选做。','评价分值：需求与设计20、核心功能25、测试与交付20、个人AI协作15、个人理解与答辩20。核心评分任务在课内完成，课后只整理与巩固。教师答案和作者测试结果不能当作学生个人证据。','## 课次安排', '|课次|主题|学时|阶段产出|','|---|---|---:|---|']
for l in lessons:md.append(f"|{l['id']}|{l['title']}|4|{l['deliverables']}|")
for l in lessons:
 md += [f"## 第{l['id']}次课 {l['title']}",f"教学主线：{l['subtitle']}。4学时，共180分钟。",'### 学习目标',l['goals'],'### 课前准备',l['preparation'],'### 重点与难点',l['pitfalls'],f"![{l['title']}的核心流程教学示意](../{l['figure']})",'图为规则和流程示意，具体运行结果以学生实际请求、回读及测试记录为准。','### 导入问题',l['question'],'### 教学过程','|学时与分钟|师生活动|形成性评价|','|---|---|---|']
 for i,period in enumerate(l['periods']):
  observation=['先写预测并指出依据，再开始操作。','教师巡查一个关键分支，学生解释自己的改动。','两人交换操作与审查角色，记录实际结果及差异。','提交本章证据；每人独立回答导入题，标记待补项。'][i]
  md.append(f'|第{i+1}学时，{i*45}—{(i+1)*45}分钟|{period}|{observation}|')
 md += ['### 网页AI任务示例',l['prompt'],'仅发送与当前任务有关的代码和脱敏错误。要求最小修改，先预测再运行；保存一次采纳、修订或拒绝的判断。','### 学生产出与评价',l['deliverables']+'。教师从“能否解释规则、改动是否正确、证据是否对应”三方面反馈，不按聊天次数或代码行数评分。','### 教学调节',l['fallback'],'### 课后巩固','整理本章任务单、最小差异和AI判断；把未通过用例写入交接，不额外堆叠新的必修功能。','### 教学反思',l['reflection'],f"配套正文：[第{l['id']}章](../{l['source']})。操作与答案入口见 [教学资源](../resources/index.html)。"]
(OUT/'完整教案.md').write_text('\n\n'.join(md)+'\n')
# HTML is offline, printable, and remains editable through the JSON/Markdown source.
css='body{max-width:1000px;margin:auto;padding:40px 24px;color:#18343c;font:17px/1.8 -apple-system,"PingFang SC","Microsoft YaHei",sans-serif}h1{font-size:36px;line-height:1.3}h2{font-size:28px;margin-top:44px}h3{font-size:20px;margin-top:28px}img{width:100%;height:auto}section{margin:42px 0;padding-top:20px;border-top:1px solid #d9e0df}table{width:100%;border-collapse:collapse;font-size:14px}td,th{border:1px solid #d9e0df;padding:12px;text-align:left;vertical-align:top}th{background:#edf3f3}nav{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:8px}a{color:#145e70}p,td{overflow-wrap:anywhere}@media(max-width:600px){nav{display:block}table{font-size:12px}td,th{padding:7px}}@media print{body{font-size:11pt;padding:0}section{break-before:page}nav{display:none}img{max-height:7cm;object-fit:contain}}'
htmlparts=[f'<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{data["title"]}</title><style>{css}</style><h1>{data["title"]}</h1><p>{intro}</p><p>{arr}</p><p><a href="../reading/index.html">图文教程</a> · <a href="完整教案.md">Markdown源稿</a> · <a href="AI编程智慧校园48学时教案.docx">Word教案</a> · <a href="AI编程智慧校园48学时教案.pdf">打印PDF</a></p><nav>']
htmlparts += [f'<a href="#lesson-{l["id"]}">{l["id"]} {l["title"]}</a>' for l in lessons];htmlparts.append('</nav>')
for l in lessons:
 htmlparts.append(f'<section id="lesson-{l["id"]}"><h2>第{l["id"]}次课 {l["title"]}</h2><p>4学时 · 180分钟 · {l["subtitle"]}</p>')
 for title,key in [('学习目标','goals'),('课前准备','preparation'),('重点与难点','pitfalls')]:htmlparts.append(f'<h3>{title}</h3><p>{html.escape(l[key])}</p>')
 htmlparts.append(f'<figure><img src="../{l["figure"]}" alt="{l["title"]}的教学流程"><figcaption>教学示意；不作为运行通过证据。</figcaption></figure><h3>导入问题</h3><p>{l["question"]}</p><h3>教学过程</h3><table><tr><th>学时</th><th>师生活动</th></tr>')
 for i,p in enumerate(l['periods']):htmlparts.append(f'<tr><td>第{i+1}学时<br>{i*45}—{(i+1)*45}分钟</td><td>{html.escape(p)}</td></tr>')
 htmlparts.append('</table>')
 for title,key in [('网页AI任务示例','prompt'),('产出与评价','deliverables'),('教学调节','fallback'),('教学反思','reflection')]:htmlparts.append(f'<h3>{title}</h3><p>{html.escape(l[key])}</p>')
 htmlparts.append('</section>')
htmlparts.append('</html>');(OUT/'index.html').write_text('\n'.join(htmlparts))
# Word: one overview page, then three pages per meeting including a full-size teaching figure.
doc=Document();section=doc.sections[0];section.page_width=Cm(21);section.page_height=Cm(29.7);section.top_margin=Cm(1.7);section.bottom_margin=Cm(1.7);section.left_margin=Cm(1.9);section.right_margin=Cm(1.9)
for style_name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=doc.styles[style_name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0);st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Heiti SC')
for st in doc.styles:
 if st.type in (1,2):
  rpr=st._element.get_or_add_rPr(); fonts=rpr.find(qn('w:rFonts'))
  if fonts is not None:
   for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']: fonts.attrib.pop(qn('w:'+attr),None)
 for border in list(st._element.iter(qn('w:pBdr'))): border.getparent().remove(border)
doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.space_after=Pt(6);doc.styles['Normal'].paragraph_format.line_spacing=1.2
doc.styles['Title'].font.size=Pt(25);doc.styles['Heading 1'].font.size=Pt(19);doc.styles['Heading 2'].font.size=Pt(13);doc.styles['Heading 2'].paragraph_format.space_before=Pt(12);doc.styles['Heading 2'].paragraph_format.space_after=Pt(5)
def para(text,style=None):return doc.add_paragraph(text,style)
def heading(text):doc.add_heading(text,2)
def table(headers, rows, widths):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,w in enumerate(widths):t.columns[i].width=Cm(w)
 for c,text in zip(t.rows[0].cells,headers):c.text=text
 for row in rows:
  cells=t.add_row().cells
  for i,(c,text) in enumerate(zip(cells,row)):c.text=text;c.width=Cm(widths[i])
 for ri,row in enumerate(t.rows):
  trPr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
  if ri==0: repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
  for c in row.cells:
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcPr=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   tcPr.append(borders);mar=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');mar.append(e)
   tcPr.append(mar)
   if ri==0:
    fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E5EDF2');tcPr.append(fill)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
    for r in p.runs:r.font.size=Pt(9.5);r.bold=(ri==0)
 return t
para(data['title'],'Title');para('48学时  12次课  教师教学用书','Subtitle');para(intro);para(arr)
heading('教学组织与评价');para('每组轮换操作与审查角色。网页AI限小任务、最小改动和人工核验；离线练习可替代网页回复。教师预置认证、连接与事务基础，学生对自己的业务改动和判断负责。AIRIOT为只读选做，占0必修学时。')
para('需求与设计20分，核心功能25分，测试与交付20分，个人AI协作15分，个人理解与答辩20分。课堂以任务单收集证据，课后整理与巩固。')
heading('课次安排')
table(['课次','主题','学时'],[[l['id'],l['title'],'4'] for l in lessons],[1.6,13,1.6])
for l in lessons:
 doc.add_page_break();doc.add_heading(f'第{l["id"]}次课  {l["title"]}',1);para(l['subtitle']+'  4学时 共180分钟')
 heading('学习目标');para(l['goals']);heading('课前准备');para(l['preparation']);heading('重点与难点');para(l['pitfalls'])
 heading('课堂板书与配图');para(f'本课第三页为可投屏、可打印的完整教学图。先让学生预测图中规则，再与代码、请求和测试结果对应。第{l["id"]}章图文教程可打开原尺寸图。')
 heading('导入问题');para(l['question']);heading('网页AI任务示例');para(l['prompt'])
 doc.add_page_break();doc.add_heading(f'第{l["id"]}次课  教学活动与评价',1)
 table(['时间','师生活动与学生任务'],[[f'第{i+1}学时\n{i*45}—{(i+1)*45}分钟',p] for i,p in enumerate(l['periods'])],[3,13.2])
 heading('形成性评价');para('先预测再操作；教师巡查关键分支并要求学生解释依据。两人每学时至少交换一次操作与审查角色。每人独立回答导入问题，课末标记已完成和待补事项。')
 heading('本次交付与反馈');para(l['deliverables']+'。评价改动、规则解释和实际证据是否对应；指出下一步最小修复，不以代码量、聊天次数或UI复刻程度评分。')
 heading('教学调节与故障替代');para(l['fallback'])
 heading('课后巩固');para('整理任务单、代码最小差异及AI判断；把未通过场景交接给下一课。核心评分项目在课内完成，不额外叠加新功能。')
 heading('教学反思');para(l['reflection'])
 para(f'配套资料  教材第{l["id"]}章、该章操作单、教师参考答案与评分要点。具体文件见完整资料包资源索引。')
 doc.add_page_break();doc.add_heading(f'第{l["id"]}次课  教学配图',1)
 para(l['subtitle']+'  课堂讨论与板书参考')
 img=ROOT/l['figure'];im=Image.open(img);w,h=im.size;scale=min(6.3/w,8.3/h)
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 pic=p.add_run().add_picture(str(img),width=Inches(w*scale),height=Inches(h*scale));pic._inline.docPr.set('descr',l['title']+'完整教学示意')
 para('教学示意；用于解释规则和流程，不代替学生实际请求、回读与测试证据。')
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=footer.add_run('AI编程智慧校园课程  ');run.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
doc.core_properties.title=data['title'];doc.core_properties.subject='48学时项目式教学';doc.core_properties.author='';doc.core_properties.last_modified_by=''
doc.save(OUT/'AI编程智慧校园48学时教案.docx')
(OUT/'README.md').write_text('# 48学时教案\n\n建议45分钟/学时，4学时×12次课，30人两人一组，课间休息另计。学校可按课表调整。\n\n- [在线阅读与打印入口](index.html)\n- [完整Markdown教案](完整教案.md)\n- [可编辑Word教案](AI编程智慧校园48学时教案.docx)\n- [打印PDF教案](AI编程智慧校园48学时教案.pdf)\n- [可供AI接力的结构化源稿](lesson-source.json)\n\n教案含逐课目标、准备、重点难点、教学图、180分钟活动、AI提示词、学生产出、评价、故障替代和反思栏。教程正文、操作单、参考工程与学生起点在完整资源包中联动维护。学校网络、Windows实机及课堂耗时由教师课前试跑；AIRIOT为0必修学时的只读选做。\n')
print('Created Markdown/HTML/DOCX for 12 lessons, 48 hours; PDF pending render review')
